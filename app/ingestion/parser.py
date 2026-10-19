"""Turn uploaded files into a list of pages with clean, offset-stable text.

PDFs keep their real page numbers, which is what makes page-level citations possible.
DOCX / TXT / MD have no physical pages, so they are split into ~3,000-character
"virtual pages" and labelled as such in the UI.
"""

from __future__ import annotations

import logging
import re
from pathlib import Path

import pymupdf

from app.schemas import Page

log = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}
VIRTUAL_PAGE_CHARS = 3000


class UnsupportedFileError(ValueError):
    pass


def clean_block(text: str) -> str:
    """Normalise one paragraph block from a PDF."""
    text = text.replace("\u00ad", "")  # soft hyphens
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)  # de-hyphenate line breaks
    text = re.sub(r"\s*\n\s*", " ", text)  # lines inside a block -> spaces
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    return text.strip()


_PAGE_NUM = re.compile(r"^(?:page\s*)?\d+(?:\s*(?:of|/)\s*\d+)?$", re.IGNORECASE)


def _boilerplate_key(block: str) -> str:
    return re.sub(r"\d+", "#", block.lower())


def parse_pdf(path: Path) -> list[Page]:
    raw_pages: list[list[str]] = []
    with pymupdf.open(path) as doc:
        for page in doc:
            # "blocks" keeps paragraphs together and headings separate, in reading order.
            blocks = page.get_text("blocks", sort=True)
            parts = [clean_block(b[4]) for b in blocks if b[6] == 0]  # b[6]==0 -> text block
            parts = [p for p in parts if p]
            if not parts:
                ocr = _try_ocr(page)
                parts = [ocr] if ocr else []
            raw_pages.append(parts)

    # Running headers/footers ("Policy X | Page 3") repeat on most pages; they add noise to
    # every chunk and can be mistaken for headings, so drop them.
    n = len(raw_pages)
    counts: dict[str, int] = {}
    for parts in raw_pages:
        for key in {_boilerplate_key(p) for p in parts}:
            counts[key] = counts.get(key, 0) + 1
    repeated = {k for k, c in counts.items() if n >= 2 and c >= max(2, 0.5 * n) and len(k) < 150}

    pages = []
    for i, parts in enumerate(raw_pages, start=1):
        kept = [p for p in parts if _boilerplate_key(p) not in repeated and not _PAGE_NUM.match(p)]
        pages.append(Page(number=i, text="\n\n".join(_join_continuations(kept))))
    return pages


_ENDS_SENTENCE = re.compile(r"[.!?:;)\]\"']$")


def _join_continuations(parts: list[str]) -> list[str]:
    """Re-join a sentence that the PDF split into separate blocks (e.g. "... 10:00 to" + "16:00.").

    A block continues the previous one when the previous block doesn't end a sentence and the
    next one starts in lower case or with a digit. Headings ("2. Annual Leave") are never merged.
    """
    from app.ingestion.chunker import looks_like_heading

    out: list[str] = []
    for part in parts:
        if (
            out
            and not _ENDS_SENTENCE.search(out[-1])
            and not looks_like_heading(out[-1])
            and (part[0].islower() or part[0].isdigit())
            and not looks_like_heading(part)
        ):
            out[-1] = f"{out[-1]} {part}"
        else:
            out.append(part)
    return out


def _try_ocr(page: pymupdf.Page) -> str:
    """Best-effort OCR for scanned pages (needs Tesseract installed); silently skipped otherwise."""
    try:
        tp = page.get_textpage_ocr(dpi=200, full=True)
        return clean_block(page.get_text("text", textpage=tp))
    except Exception:  # noqa: BLE001 - OCR is optional
        log.info("Page %s has no text layer and OCR is unavailable", page.number + 1)
        return ""


def _virtual_pages(full_text: str) -> list[Page]:
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", full_text) if p.strip()]
    pages, buf, size = [], [], 0
    for para in paragraphs:
        if buf and size + len(para) > VIRTUAL_PAGE_CHARS:
            pages.append(Page(number=len(pages) + 1, text="\n\n".join(buf)))
            buf, size = [], 0
        buf.append(para)
        size += len(para) + 2
    if buf:
        pages.append(Page(number=len(pages) + 1, text="\n\n".join(buf)))
    return pages


def parse_docx(path: Path) -> list[Page]:
    import docx  # python-docx

    d = docx.Document(str(path))
    paras = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    for table in d.tables:  # flatten tables row-by-row
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                paras.append(" | ".join(cells))
    return _virtual_pages("\n\n".join(paras))


def parse_text(path: Path) -> list[Page]:
    return _virtual_pages(path.read_text(encoding="utf-8", errors="replace"))


def parse_file(path: Path) -> list[Page]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        pages = parse_pdf(path)
    elif ext == ".docx":
        pages = parse_docx(path)
    elif ext in {".txt", ".md"}:
        pages = parse_text(path)
    else:
        raise UnsupportedFileError("This file type isn't supported. Upload a PDF, Word (.docx), text or Markdown file.")
    if not any(p.text.strip() for p in pages):
        raise UnsupportedFileError(
            "No readable text was found. The file may be a scanned image, which needs OCR to be read."
        )
    return pages
